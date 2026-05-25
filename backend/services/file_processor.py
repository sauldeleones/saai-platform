import io
import os
from pathlib import Path

from models.entrega import TipoArchivo

# Extensiones reconocidas por tipo
EXTENSIONES = {
    TipoArchivo.document: {".pdf", ".docx", ".doc", ".txt", ".md"},
    TipoArchivo.code: {".py", ".js", ".ts", ".java", ".c", ".cpp", ".cs", ".go",
                       ".rb", ".php", ".swift", ".kt", ".r", ".sql", ".html", ".css"},
    TipoArchivo.presentation: {".pptx", ".ppt"},
    TipoArchivo.image: {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"},
}


def detectar_tipo(nombre_archivo: str) -> TipoArchivo:
    ext = Path(nombre_archivo).suffix.lower()
    for tipo, extensiones in EXTENSIONES.items():
        if ext in extensiones:
            return tipo
    return TipoArchivo.document  # fallback


def extraer_texto(ruta_archivo: str, tipo_archivo: TipoArchivo) -> str:
    """Punto de entrada principal — delega según el tipo de archivo."""
    try:
        if tipo_archivo == TipoArchivo.document:
            return _extraer_documento(ruta_archivo)
        elif tipo_archivo == TipoArchivo.code:
            return _extraer_codigo(ruta_archivo)
        elif tipo_archivo == TipoArchivo.presentation:
            return _extraer_presentacion(ruta_archivo)
        elif tipo_archivo == TipoArchivo.image:
            return _extraer_imagen(ruta_archivo)
    except Exception as e:
        return f"[Error extrayendo texto: {str(e)}]"
    return ""


def _extraer_documento(ruta: str) -> str:
    ext = Path(ruta).suffix.lower()

    if ext == ".pdf":
        return _extraer_pdf(ruta)
    elif ext in (".docx", ".doc"):
        return _extraer_docx(ruta)
    else:
        # .txt, .md y otros: leer directo
        with open(ruta, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()


def _extraer_pdf(ruta: str) -> str:
    import PyPDF2
    texto = []
    with open(ruta, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for pagina in reader.pages:
            contenido = pagina.extract_text()
            if contenido:
                texto.append(contenido)
    return "\n".join(texto)


def _extraer_docx(ruta: str) -> str:
    from docx import Document
    doc = Document(ruta)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extraer_codigo(ruta: str) -> str:
    with open(ruta, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extraer_presentacion(ruta: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(ruta)
        texto = []
        for i, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texto.append(f"[Diapositiva {i}] {shape.text.strip()}")
        return "\n".join(texto)
    except ImportError:
        return "[Instala python-pptx para procesar presentaciones: pip install python-pptx]"


def _extraer_imagen(ruta: str) -> str:
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(ruta)
        return pytesseract.image_to_string(img, lang="spa+eng")
    except Exception as e:
        # Tesseract puede no estar instalado — no es un error fatal
        return f"[OCR no disponible: {str(e)}. Instala tesseract: brew install tesseract]"
